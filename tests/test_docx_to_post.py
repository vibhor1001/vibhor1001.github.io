"""
tools/docx-to-post.py — a Word document uploaded in the editor becomes a draft.

HOW TO RUN   pip install -r tools/requirements.txt python-docx
             python3 -m unittest tests/test_docx_to_post.py

python-docx is only for these tests: it writes the Word documents they feed
in, shaped like marketing's (label lines, a Heading 1 title, a table, a
captioned picture).
"""
import importlib.util
import io
import shutil
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path

from docx import Document
from docx.shared import Inches
from PIL import Image

REPO = Path(__file__).resolve().parent.parent


def load(name, file):
    spec = importlib.util.spec_from_file_location(name, REPO / "tools" / file)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod  # dataclasses look the module up by name
    spec.loader.exec_module(mod)
    return mod


dtp = load("docx_to_post", "docx-to-post.py")
bb = load("build_blog_for_docx_tests", "build-blog.py")


def rendered(markdown_text: str) -> str:
    """What the site shows for this Markdown (the blog builder's own renderer)."""
    return bb.to_html(markdown_text)


def png(width=2400, height=1350, colour=(200, 30, 90)) -> io.BytesIO:
    buf = io.BytesIO()
    Image.new("RGB", (width, height), colour).save(buf, "PNG")
    buf.seek(0)
    return buf


def word_post(path: Path, *, labels=True, url="/my-first-post/", title="My First Post: A Guide",
              picture=True, caption="(A chart of bookings by channel)", extra=None, title_style=False,
              picture_size=(2400, 1350), rows=None):
    doc = Document()
    if labels:
        p = doc.add_paragraph()
        p.add_run("Meta Title:").bold = True
        p.add_run(" My First Post | PropertyFlow")
        if url is not None:
            p = doc.add_paragraph()
            p.add_run("Recommended URL").bold = True
            p.add_run(f": {url}")
        p = doc.add_paragraph()
        p.add_run("Primary Keyword: ").bold = True
        p.add_run("first post")
    if title:
        doc.add_heading(title, level=0 if title_style else 1)   # level 0 = Word's "Title" style
    if labels:
        p = doc.add_paragraph()
        p.add_run("Meta Description: ").bold = True
        p.add_run("A short description for Google.")
    doc.add_paragraph("Opening paragraph with some ").add_run("bold words").bold = True
    doc.add_heading("A section", level=2)
    doc.add_paragraph("Pricing: this sentence has a colon but is ordinary text, not a label.")
    rows = rows or [("Task", "Tool"), ("Calendars", "Channel Manager"), ("Rates", "Cognito")]
    table = doc.add_table(rows=len(rows), cols=2)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    if picture:
        doc.add_picture(png(*picture_size), width=Inches(5))
        if caption:
            doc.add_paragraph().add_run(caption).italic = True
    for line in extra or []:
        doc.add_paragraph(line)
    doc.save(path)


class DocxToPost(unittest.TestCase):
    def setUp(self):
        self.root = Path(tempfile.mkdtemp())
        (self.root / "content" / "blog").mkdir(parents=True)
        (self.root / "images" / "blog").mkdir(parents=True)
        self.saved = (dtp.ROOT, dtp.INBOX, dtp.POSTS, dtp.IMAGES)
        dtp.ROOT, dtp.INBOX, dtp.POSTS, dtp.IMAGES = (
            self.root, self.root / "images" / "blog", self.root / "content" / "blog",
            self.root / "images" / "blog" / "word")

    def tearDown(self):
        dtp.ROOT, dtp.INBOX, dtp.POSTS, dtp.IMAGES = self.saved
        shutil.rmtree(self.root)

    def upload(self, name="my-first-post.docx", **kw) -> Path:
        path = self.root / "images" / "blog" / name
        word_post(path, **kw)
        return path

    def run_script(self):
        old = sys.argv
        sys.argv = ["docx-to-post.py"]
        out = io.StringIO()
        try:
            with redirect_stdout(out):
                code = dtp.main()
        finally:
            sys.argv = old
        return code, out.getvalue()

    def post(self, slug="my-first-post"):
        import yaml
        text = (self.root / "content" / "blog" / f"{slug}.md").read_text(encoding="utf-8")
        _, front, body = text.split("---\n", 2)
        return yaml.safe_load(front), body

    # ─── The fields come from the label lines ────────────────────────────────

    def test_labels_fill_the_fields_and_leave_the_text(self):
        doc = self.upload()
        code, _ = self.run_script()
        self.assertEqual(code, 0)
        meta, body = self.post()
        self.assertEqual(meta["title"], "My First Post: A Guide")
        self.assertEqual(meta["url"], "my-first-post")
        self.assertEqual(meta["meta_title"], "My First Post | PropertyFlow")
        self.assertEqual(meta["description"], "A short description for Google.")
        self.assertIs(meta["draft"], True)
        for label in ("Meta Title", "Recommended URL", "Primary Keyword", "Meta Description", "first post"):
            self.assertNotIn(label, body)
        self.assertNotIn("# My First Post", body)             # the title is not repeated in the text
        self.assertIn("Pricing: this sentence has a colon", body)  # ordinary text with a colon stays
        self.assertFalse(doc.exists(), "the upload is removed once it is a draft")

    def test_text_headings_bold_and_table_carry_over(self):
        self.upload()
        self.run_script()
        _, body = self.post()
        self.assertIn("Opening paragraph with some **bold words**", body)
        self.assertIn("## A section", body)
        self.assertIn("| Task | Tool |\n| --- | --- |\n| Calendars | Channel Manager |", body)

    def test_picture_becomes_a_small_webp_with_its_caption_as_description(self):
        self.upload()
        self.run_script()
        meta, body = self.post()
        webp = self.root / "images" / "blog" / "word" / "my-first-post" / "1.webp"
        self.assertTrue(webp.exists())
        with Image.open(webp) as img:
            self.assertEqual(img.format, "WEBP")
            self.assertEqual(img.width, 1600)
        self.assertIn("![A chart of bookings by channel](/images/blog/word/my-first-post/1.webp)", body)
        self.assertEqual(meta["thumbnail"], "/images/blog/word/my-first-post/1.webp")
        self.assertEqual(meta["thumbnail_alt"], "A chart of bookings by channel")

    def test_without_labels_the_title_gives_the_web_address(self):
        self.upload(name="draft.docx", labels=False, title="Edinburgh Short Let Licences Explained")
        self.run_script()
        meta, _ = self.post("edinburgh-short-let-licences-explained")
        self.assertEqual(meta["description"], "")
        self.assertEqual(meta["meta_title"], "")

    def test_a_full_address_or_blog_path_in_the_url_line_still_gives_the_last_part(self):
        self.upload(url="https://www.propertyflow.uk/blog/Pricing-Tips/")
        self.run_script()
        self.assertTrue((self.root / "content" / "blog" / "pricing-tips.md").exists())

    def test_no_picture_leaves_the_card_image_for_marketing(self):
        self.upload(picture=False)
        code, _ = self.run_script()
        self.assertEqual(code, 0)
        meta, _ = self.post()
        self.assertEqual(meta["thumbnail"], "")

    # ─── Re-uploads and clashes ──────────────────────────────────────────────

    def test_re_uploading_replaces_the_draft_and_its_pictures(self):
        self.upload(extra=["First version."])
        self.run_script()
        self.upload(extra=["Second version."], picture=False)
        self.run_script()
        _, body = self.post()
        self.assertIn("Second version.", body)
        self.assertNotIn("First version.", body)
        self.assertFalse((self.root / "images" / "blog" / "word" / "my-first-post" / "1.webp").exists())

    def test_other_posts_pictures_are_never_touched(self):
        """Re-uploading "london" must not delete london-2024.webp that a live post uses."""
        blog = self.root / "images" / "blog"
        theirs = [blog / "my-first-post-2024.webp", blog / "my-first-post-1.webp",
                  blog / "word" / "my-first-post-guide" / "1.webp"]
        for f in theirs:
            f.parent.mkdir(parents=True, exist_ok=True)
            f.write_bytes(b"x")
        self.upload()
        self.run_script()
        self.upload()
        self.run_script()
        for f in theirs:
            self.assertTrue(f.exists(), f.name)

    def test_a_live_post_without_a_draft_line_or_with_broken_details_is_never_overwritten(self):
        for text in ("---\ntitle: Live\n---\nLive text.\n", "---\ntitle: [unclosed\n---\nLive text.\n"):
            live = self.root / "content" / "blog" / "my-first-post.md"
            live.write_text(text, encoding="utf-8")
            self.upload()
            self.run_script()
            self.assertIn("Live text.", live.read_text())
            (self.root / "content" / "blog" / "my-first-post-2.md").unlink()

    def test_a_live_post_is_never_overwritten(self):
        live = self.root / "content" / "blog" / "my-first-post.md"
        live.write_text("---\ntitle: Live\ndraft: false\n---\nLive text.\n", encoding="utf-8")
        self.upload()
        self.run_script()
        self.assertIn("Live text.", live.read_text())
        meta, _ = self.post("my-first-post-2")
        self.assertIs(meta["draft"], True)

    def test_the_list_pages_addresses_are_not_taken(self):
        self.upload(url="/oldest/")
        self.run_script()
        self.assertTrue((self.root / "content" / "blog" / "oldest-post.md").exists())
        self.assertFalse((self.root / "content" / "blog" / "oldest.md").exists())

    # ─── Labels only at the top ──────────────────────────────────────────────

    def test_a_label_like_line_in_the_text_stays_text(self):
        self.upload(labels=False, extra=["Description: lead with the problem the host has.",
                                         "URL: https://www.gov.uk/short-term-lets is where the rules live.",
                                         "Author: Jane"])
        self.run_script()
        meta, body = self.post("my-first-post-a-guide")
        self.assertEqual(meta["description"], "")
        for line in ("Description: lead with the problem", "URL: https://www.gov.uk", "Author: Jane"):
            self.assertIn(line, body)

    def test_labels_on_one_line_with_shift_enter_and_a_value_on_the_next_line(self):
        doc = Document()
        p = doc.add_paragraph("Meta Title: Pricing Tips | PropertyFlow")
        p.add_run().add_break()
        p.add_run("Recommended URL: /pricing-tips/")
        doc.add_heading("Pricing Tips", level=1)
        doc.add_paragraph("Meta Description:")
        doc.add_paragraph("How to set nightly rates.")
        doc.add_paragraph("The post starts here.")
        doc.save(self.root / "images" / "blog" / "tips.docx")
        self.run_script()
        meta, body = self.post("pricing-tips")
        self.assertEqual(meta["meta_title"], "Pricing Tips | PropertyFlow")
        self.assertEqual(meta["description"], "How to set nightly rates.")
        self.assertNotIn("nightly rates", body)
        self.assertIn("The post starts here.", body)

    def test_word_title_style_is_the_title(self):
        self.upload(title_style=True)
        self.run_script()
        meta, body = self.post()
        self.assertEqual(meta["title"], "My First Post: A Guide")
        self.assertNotIn("My First Post: A Guide", body)

    def test_no_heading_uses_the_google_title_without_the_brand(self):
        self.upload(title=None)
        self.run_script()
        meta, _ = self.post()
        self.assertEqual(meta["title"], "My First Post")

    def test_a_title_with_quotes_and_colons_comes_back_exactly(self):
        title = 'The "Airbnb" loophole: C:\\path & more'
        self.upload(labels=False, title=title)
        self.run_script()
        [md] = list((self.root / "content" / "blog").glob("*.md"))
        import yaml
        self.assertEqual(yaml.safe_load(md.read_text().split("---\n")[1])["title"], title)

    # ─── Words stay words ────────────────────────────────────────────────────

    def test_characters_typed_in_word_show_as_typed_and_never_become_code(self):
        self.upload(extra=["Hi <guest name>, welcome.", "<script>alert(1)</script>",
                           "#1 mistake hosts make.", "£29* a month, trial*.", "2026. A year of change.",
                           "- not a list", "Fees & charges"])
        self.run_script()
        _, body = self.post()
        page = rendered(body)
        self.assertNotIn("<script>", page)
        self.assertIn("&lt;script&gt;alert(1)&lt;/script&gt;", page)
        self.assertIn("Hi &lt;guest name&gt;, welcome.", page)
        self.assertIn("#1 mistake hosts make.", page)
        self.assertIn("£29* a month, trial*.", page)
        self.assertIn("2026. A year of change.", page)
        self.assertIn("- not a list", page)
        self.assertIn("Fees &amp; charges", page)
        self.assertNotIn("<ol", page)
        self.assertNotIn("<em>", page.split("£29")[1].split("</p>")[0])

    def test_only_web_email_and_phone_links_survive(self):
        md = dtp.to_markdown('<p><a href="javascript:alert(1)">click</a> <a href="https://x.uk/a">site</a> '
                             '<a href="mailto:a@b.uk">mail</a> <a href="data:text/html,x">data</a></p>')
        self.assertNotIn("javascript", md)
        self.assertNotIn("data:", md)
        self.assertIn("[site](https://x.uk/a)", md)
        self.assertIn("[mail](mailto:a@b.uk)", md)
        self.assertIn("click", md)

    def test_a_pipe_in_a_table_cell_keeps_the_table_whole(self):
        self.upload(rows=[("Plan", "Price"), ("Basic | Pro", "£10")])
        self.run_script()
        _, body = self.post()
        page = rendered(body)
        self.assertIn("<td>Basic | Pro</td>", page)
        self.assertIn("<td>£10</td>", page)

    def test_nested_lists_stay_nested_on_the_page(self):
        md = dtp.to_markdown("<ul><li>Channels<ul><li>Airbnb</li><li>Booking.com</li></ul></li><li>Pricing</li></ul>"
                             "<ol><li>First<ol><li>Inner</li></ol></li></ol>")
        page = rendered(md)
        self.assertRegex(page, r"<li>Channels\s*<ul>\s*<li>Airbnb</li>")
        self.assertRegex(page, r"<li>First\s*<ol>\s*<li>Inner</li>")

    # ─── Pictures ────────────────────────────────────────────────────────────

    def test_a_very_tall_picture_is_shrunk_not_a_crash(self):
        self.upload(picture_size=(800, 17000))
        code, _ = self.run_script()
        self.assertEqual(code, 0)
        with Image.open(self.root / "images" / "blog" / "word" / "my-first-post" / "1.webp") as img:
            self.assertLessEqual(img.height, 16383)

    def test_a_picture_that_cannot_be_converted_is_left_out_with_a_note_marketing_will_see(self):
        body, saved, dropped = dtp.save_images('<p><img alt="" src="@@IMG0@@" /></p>', [b"not an image"], "s", "T")
        self.assertEqual((saved, dropped), ([], 1))
        self.assertNotIn("<img", body)
        real = dtp.read_docx
        dtp.read_docx = lambda path: ('<h1>T</h1><p>Text.</p><p><img alt="" src="@@IMG0@@" /></p>', [b"junk"])
        try:
            (self.root / "images" / "blog" / "t.docx").write_bytes(b"x")
            code, _ = self.run_script()
        finally:
            dtp.read_docx = real
        self.assertEqual(code, 0)
        _, body = self.post("t")
        self.assertIn("Note from the Word upload:** 1 picture", body)

    def test_word_s_own_picture_description_gives_way_to_the_caption_or_title(self):
        html_in = ('<p><img alt="A screenshot of a computer\n\nAI-generated content may be incorrect." src="@@IMG0@@" /></p>'
                   '<p><em>(Bookings by channel)</em></p><p><img alt="" src="@@IMG1@@" /></p>')
        body, saved, _ = dtp.save_images(html_in, [png().getvalue(), png().getvalue()], "s", "The title")
        self.assertEqual([alt for _, alt in saved], ["Bookings by channel", "The title"])

    # ─── Pictures uploaded straight into the editor ──────────────────────────

    def test_a_big_picture_uploaded_in_the_editor_is_shrunk_in_place(self):
        blog = self.root / "images" / "blog"
        # Smooth pictures: light enough once shrunk to stay as they are (a heavy
        # photo-like PNG is converted instead; see the WebP test below).
        smooth = Image.linear_gradient("L").resize((3200, 1800)).convert("RGB")
        smooth.save(blog / "hero.jpg", "JPEG", quality=100)
        smooth.resize((3000, 1500)).save(blog / "chart.png", "PNG")
        before = {n: (blog / n).stat().st_size for n in ("hero.jpg", "chart.png")}
        code, out = self.run_script()
        self.assertEqual(code, 0)
        for name in ("hero.jpg", "chart.png"):
            with Image.open(blog / name) as img:
                self.assertEqual(img.width, 1600, name)          # same name, same place: posts still find it
            self.assertLess((blog / name).stat().st_size, before[name], name)
        self.assertIn("shrank hero.jpg", out)

    def test_a_small_picture_is_left_exactly_as_it_is_and_not_recompressed_each_run(self):
        blog = self.root / "images" / "blog"
        (blog / "small.png").write_bytes(png(800, 450).getvalue())
        Image.effect_noise((3200, 1800), 80).convert("RGB").save(blog / "big.jpg", "JPEG", quality=100)
        small = (blog / "small.png").read_bytes()
        self.run_script()
        once = (blog / "big.jpg").read_bytes()
        self.run_script()
        self.assertEqual((blog / "small.png").read_bytes(), small)
        self.assertEqual((blog / "big.jpg").read_bytes(), once)   # second run: no further loss

    def test_a_phone_photo_stays_the_right_way_up(self):
        blog = self.root / "images" / "blog"
        img = Image.effect_noise((3000, 2000), 80).convert("RGB")      # stored landscape …
        exif = Image.Exif()
        exif[0x0112] = 6                                                # … "rotate 90°": shows as portrait
        img.save(blog / "phone.jpg", "JPEG", quality=95, exif=exif)
        self.run_script()
        with Image.open(blog / "phone.jpg") as out:
            self.assertLess(out.width, out.height)                     # still portrait
            self.assertLessEqual(out.width, 1600)

    def test_animations_and_colour_profiles_survive(self):
        blog = self.root / "images" / "blog"
        frames = [Image.effect_noise((1800, 600), 60 + i * 10).convert("RGB") for i in range(4)]
        frames[0].save(blog / "anim.webp", "WEBP", save_all=True, append_images=frames[1:], duration=200)
        anim = (blog / "anim.webp").read_bytes()
        from PIL import ImageCms
        icc = ImageCms.ImageCmsProfile(ImageCms.createProfile("sRGB")).tobytes()
        Image.effect_noise((3000, 1500), 80).convert("RGB").save(blog / "photo.jpg", "JPEG", quality=100, icc_profile=icc)
        self.run_script()
        self.assertEqual((blog / "anim.webp").read_bytes(), anim)
        with Image.open(blog / "photo.jpg") as out:
            self.assertEqual(out.width, 1600)
            self.assertEqual(out.info.get("icc_profile"), icc)

    def test_an_upper_case_name_is_shrunk_too(self):
        blog = self.root / "images" / "blog"
        Image.effect_noise((3200, 400), 60).convert("RGB").save(blog / "BANNER.JPG", "JPEG", quality=30)
        self.run_script()
        with Image.open(blog / "BANNER.JPG") as out:
            self.assertEqual(out.width, 1600)

    def test_a_wide_picture_that_would_not_get_lighter_is_left_alone(self):
        blog = self.root / "images" / "blog"
        # Very low quality and wide: a 1600px copy at normal quality is heavier.
        Image.effect_noise((2400, 300), 60).convert("RGB").save(blog / "low.jpg", "JPEG", quality=10)
        before = (blog / "low.jpg").read_bytes()
        self.run_script()
        self.assertEqual((blog / "low.jpg").read_bytes(), before)

    def test_a_heavy_picture_under_1600_wide_is_recompressed(self):
        blog = self.root / "images" / "blog"
        Image.effect_noise((1500, 1000), 80).convert("RGB").save(blog / "heavy.jpg", "JPEG", quality=100)
        before = (blog / "heavy.jpg").stat().st_size
        self.assertGreater(before, dtp.BIG_BYTES)
        self.run_script()
        self.assertLess((blog / "heavy.jpg").stat().st_size, before)

    def test_a_cmyk_photo_stays_readable(self):
        blog = self.root / "images" / "blog"
        Image.effect_noise((3000, 1500), 80).convert("CMYK").save(blog / "print.jpg", "JPEG", quality=100)
        self.run_script()
        with Image.open(blog / "print.jpg") as out:
            out.load()
            self.assertEqual(out.width, 1600)

    def test_a_heavy_photo_saved_as_png_becomes_webp_and_the_post_follows_it(self):
        blog, posts = self.root / "images" / "blog", self.root / "content" / "blog"
        # Photo-like PNG over the limit even at 1600px (Post 1's card image was 1.7 MB).
        Image.effect_noise((1672, 941), 90).convert("RGB").save(blog / "card.png", "PNG")
        (posts / "live.md").write_text("---\ntitle: L\ndraft: false\nthumbnail: /images/blog/card.png\n---\n![x](/images/blog/card.png)\n")
        (posts / "other.md").write_text("---\ntitle: O\n---\n![y](/images/blog/other.png)\n")
        self.run_script()
        self.assertFalse((blog / "card.png").exists())
        with Image.open(blog / "card.webp") as out:
            self.assertEqual(out.format, "WEBP")
            self.assertLessEqual(out.width, 1600)
        live = (posts / "live.md").read_text()
        self.assertEqual(live.count("/images/blog/card.webp"), 2)
        self.assertNotIn("card.png", live)
        self.assertIn("/images/blog/other.png", (posts / "other.md").read_text())

    def test_webp_conversion_only_rewrites_whole_names(self):
        blog, posts = self.root / "images" / "blog", self.root / "content" / "blog"
        Image.effect_noise((1672, 941), 90).convert("RGB").save(blog / "a.png", "PNG")
        (posts / "p.md").write_text("---\ntitle: P\nthumbnail: /images/blog/a.png\n---\n"
                                    "![x](/images/blog/a.png) ![y](/images/blog/a.png.bak) ![z](/images/blog/xa.png)\n")
        self.run_script()
        text = (posts / "p.md").read_text()
        self.assertIn("thumbnail: /images/blog/a.webp", text)
        self.assertIn("![x](/images/blog/a.webp)", text)
        self.assertIn("/images/blog/a.png.bak", text)
        self.assertIn("/images/blog/xa.png", text)

    def test_cmyk_and_16_bit_pictures_are_not_converted_to_webp(self):
        blog = self.root / "images" / "blog"
        Image.effect_noise((1600, 1000), 90).convert("CMYK").save(blog / "print.jpg", "JPEG", quality=100)
        Image.effect_noise((1600, 1000), 90).convert("I;16").save(blog / "grey.png", "PNG")
        self.run_script()
        self.assertTrue((blog / "print.jpg").exists())
        self.assertTrue((blog / "grey.png").exists())
        self.assertFalse((blog / "print.webp").exists())
        self.assertFalse((blog / "grey.webp").exists())

    def test_no_webp_conversion_when_that_name_is_taken(self):
        blog = self.root / "images" / "blog"
        Image.effect_noise((1672, 941), 90).convert("RGB").save(blog / "card.png", "PNG")
        (blog / "card.webp").write_bytes(b"someone else's")
        self.run_script()
        self.assertTrue((blog / "card.png").exists())
        self.assertEqual((blog / "card.webp").read_bytes(), b"someone else's")

    def test_the_word_folder_and_gifs_are_not_touched(self):
        word = self.root / "images" / "blog" / "word" / "p"
        word.mkdir(parents=True)
        Image.effect_noise((3200, 1800), 80).convert("RGB").save(word / "1.webp", "WEBP", quality=100)
        gif = self.root / "images" / "blog" / "anim.gif"
        Image.effect_noise((2400, 900), 80).convert("P").save(gif)   # wide and heavy: would be shrunk if not skipped
        before = ((word / "1.webp").read_bytes(), gif.read_bytes())
        self.run_script()
        self.assertEqual(((word / "1.webp").read_bytes(), gif.read_bytes()), before)

    # ─── Web address changed on a draft ──────────────────────────────────────

    def test_a_draft_whose_web_address_was_changed_is_renamed_but_a_live_post_is_not(self):
        posts = self.root / "content" / "blog"
        (posts / "old-name.md").write_text("---\ntitle: D\nurl: new-name\ndraft: true\n---\nText\n")
        (posts / "live-name.md").write_text("---\ntitle: L\nurl: other-name\ndraft: false\n---\nText\n")
        self.run_script()
        self.assertTrue((posts / "new-name.md").exists())
        self.assertFalse((posts / "old-name.md").exists())
        self.assertTrue((posts / "live-name.md").exists())
        self.assertFalse((posts / "other-name.md").exists())

    # ─── A document that cannot be read ──────────────────────────────────────

    def test_a_file_that_is_not_a_word_document_is_left_and_reported(self):
        bad = self.root / "images" / "blog" / "notes.docx"
        bad.write_bytes(b"this is not a zip file")
        self.upload()
        code, out = self.run_script()
        self.assertEqual(code, 2)                  # fails the Action → GitHub emails the owner
        self.assertTrue(bad.exists())
        self.assertIn("::warning::docx-to-post: notes.docx could not be read", out)
        self.assertTrue((self.root / "content" / "blog" / "my-first-post.md").exists(), "the good one still converts")

    def test_an_upload_that_is_not_a_docx_is_reported(self):
        (self.root / "images" / "blog" / "old-post.doc").write_bytes(b"x")
        code, out = self.run_script()
        self.assertEqual(code, 2)
        self.assertIn("old-post.doc is not a Word .docx file", out)

    def test_a_zip_bomb_is_refused(self):
        path = self.root / "images" / "blog" / "bomb.docx"
        import zipfile
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("word/document.xml", b"0" * (dtp.MAX_UNZIPPED + 1))
        code, out = self.run_script()
        self.assertEqual(code, 2)
        self.assertIn("too large once unpacked", out)
        self.assertTrue(path.exists())

    def test_an_unexpected_error_in_one_document_does_not_block_the_next(self):
        self.upload(name="a.docx")
        self.upload(name="b.docx", url="/second/")
        real = dtp.save_images
        calls = []
        def flaky(body, images, slug, alt):
            calls.append(slug)
            if len(calls) == 1:
                raise RuntimeError("boom")
            return real(body, images, slug, alt)
        dtp.save_images = flaky
        try:
            code, out = self.run_script()
        finally:
            dtp.save_images = real
        self.assertEqual(code, 2)
        self.assertIn("a.docx could not be converted (RuntimeError: boom)", out)
        self.assertTrue((self.root / "content" / "blog" / "second.md").exists())

    def test_an_empty_document_is_left_and_reported(self):
        path = self.root / "images" / "blog" / "empty.docx"
        Document().save(path)
        code, out = self.run_script()
        self.assertEqual(code, 2)
        self.assertTrue(path.exists())
        self.assertIn("has no text", out)

    def test_the_draft_is_read_by_the_blog_builder_without_complaint(self):
        """A draft with no card image or description must not fail the build."""
        spec_b = importlib.util.spec_from_file_location("build_blog_t", REPO / "tools" / "build-blog.py")
        bb = importlib.util.module_from_spec(spec_b)
        sys.modules["build_blog_t"] = bb
        spec_b.loader.exec_module(bb)
        self.upload(labels=False, picture=False)
        self.run_script()
        [md] = list((self.root / "content" / "blog").glob("*.md"))
        post = bb.parse_post(md)
        self.assertTrue(post.draft)


if __name__ == "__main__":
    unittest.main()
